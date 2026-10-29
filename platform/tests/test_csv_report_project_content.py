"""Cross-library report: independent totals, real workbook and meaningful plot."""
import contextlib
import copy
import io
import json
from pathlib import Path
import sys
import pytest

ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT))
import build_tasks
from test_project_courses import client,srv,register,solve_topic


def quest():
    section=build_tasks.SRC.read_text(encoding='utf-8').split('\n## Мини-квесты на своём компьютере\n',1)[1]
    matches=[q for q in json.loads(build_tasks.first_fence(section)[1]) if q['id']=='practice_csv_report']
    assert len(matches)==1,'The independent CSV → Excel/PNG project is missing'
    return matches[0]


def test_three_connected_stages_have_explicit_library_prerequisites():
    course=quest()
    assert [s['id'] for s in course['stages']]==['audit_csv','write_excel','plot_report']
    assert {49,52,53}<=set(course['unlock_topics'])
    assert 56 in course['stages'][1]['requirements']['topics']
    assert 60 in course['stages'][2]['requirements']['topics']


def checked(source,check):
    out=io.StringIO()
    with contextlib.redirect_stdout(out):
        exec(source+'\n'+check,{})
    assert out.getvalue().strip()=='OK'


@pytest.mark.parametrize('stage_index',[0,1,2])
@pytest.mark.parametrize('variant',['reference','alternative'])
def test_equivalent_implementations_pass_real_result_checks(stage_index,variant,tmp_path,monkeypatch):
    monkeypatch.chdir(tmp_path)
    stage=quest()['stages'][stage_index];source=stage['reference_code']
    if variant=='alternative' and stage_index==0:
        start=source.index('    frame = pd.DataFrame')
        source=source[:start]+'''    totals = {}
    for _, item, qty, cents in clean:
        record = totals.setdefault(item, {'item':item,'qty':0,'amount_cents':0})
        record['qty'] += qty
        record['amount_cents'] += qty*cents
    groups = [totals[name] for name in sorted(totals)]
    return {'rows':len(clean),'qty':sum(row[2] for row in clean),
            'amount_cents':sum(row[2]*row[3] for row in clean),'groups':groups}
'''
    elif variant=='alternative' and stage_index==1:
        source=source.replace("with open(path, 'xb') as output:\n        book.save(output)",
            "from io import BytesIO\n    payload = BytesIO()\n    book.save(payload)\n    with open(path, 'xb') as output:\n        output.write(payload.getvalue())")
    elif variant=='alternative':
        source=source.replace("ax.bar([row['item'] for row in rows], [row['amount_cents'] for row in rows])",
            "ax.bar(range(len(rows)), [row['amount_cents'] for row in rows], color='teal', width=0.6)")
    checked(source,stage['check']['append'])


MUTATIONS=[
    (0,'or record_id in ids',''),
    (0,"if len(row) != 4:","if len(row) > 4:"),
    (0,"frame['qty'] * frame['unit_cents']","frame['qty']"),
    (0,"as_index=False, sort=True","as_index=False, sort=False"),
    (0,"int(frame['qty'].sum())","float(frame['qty'].sum())"),
    (0,"if len(rows) - 1 > 500:","if len(rows) - 1 > 501:"),
    (0,"qty, cents = int(qty), int(cents)","qty, cents = (int(cents),int(qty)) if rows[0][0]=='item' else (int(qty),int(cents))"),
    (1,"sheet.cell(sheet.max_row, 1).data_type = 's'","pass"),
    (1,"open(path, 'xb')","open(path, 'wb')"),
    (1,"group['qty'], group['amount_cents']","str(group['qty']), group['amount_cents']"),
    (1,"['Количество', summary['qty']]","['Количество', summary['rows']]"),
    (1,"book.close()","book.close()\n    summary['groups'].clear()"),
    (2,"-row['amount_cents']","row['amount_cents']"),
    (2,"[:5]","[:4]"),
    (2,"[row['amount_cents'] for row in rows]","[row['qty'] for row in rows]"),
    (2,"ax.set_ylabel('Сумма, коп.')","ax.set_ylabel('Сумма, руб.')"),
    (2,"ax.text(0.5, 0.5, 'Нет данных', transform=ax.transAxes, ha='center')","pass"),
    (2,"return figure","summary['groups'].reverse()\n    return figure"),
]


@pytest.mark.parametrize('index,old,new',MUTATIONS)
def test_checks_reject_independent_report_defects(index,old,new,tmp_path,monkeypatch):
    monkeypatch.chdir(tmp_path)
    stage=quest()['stages'][index]
    source=stage['reference_code'];assert old in source
    with pytest.raises((AssertionError,ValueError,IndexError)):
        checked(source.replace(old,new),stage['check']['append'])


@pytest.mark.parametrize('index',[0,1,2])
@pytest.mark.parametrize('code_kind',['starter_code','print_only'])
def test_empty_or_print_only_is_not_a_report(index,code_kind):
    stage=quest()['stages'][index]
    source=stage.get(code_kind,"print('OK')")
    with pytest.raises((AssertionError,TypeError,NameError,AttributeError,OSError)):
        checked(source,stage['check']['append'])


def test_sample_connects_all_three_stages_and_preserves_inputs(tmp_path):
    from openpyxl import load_workbook
    namespace={}
    for stage in quest()['stages']:exec(stage['reference_code'],namespace)
    text='record_id,item,qty,unit_cents\n001,Чай,2,1250\n002,Кофе,1,800\n003,Чай,3,200\n'
    summary=namespace['audit_sales'](text)
    assert (summary['rows'],summary['qty'],summary['amount_cents'])==(3,6,3900)
    original=copy.deepcopy(summary)
    target=tmp_path/'sales.xlsx';namespace['write_sales_excel'](summary,target)
    book=load_workbook(target,data_only=False)
    assert list(book['Сводка'].values)[1:]==[('Кофе',1,800),('Чай',5,3100)]
    assert book['Контроль']['B4'].value==3900
    book.close()
    fig=namespace['plot_sales'](summary)
    assert [p.get_height() for p in fig.axes[0].patches]==[3100,800]
    fig.savefig(tmp_path/'sales.png');fig.clear()
    assert (tmp_path/'sales.png').read_bytes().startswith(b'\x89PNG\r\n\x1a\n')
    assert summary==original


def test_generated_project_matches_source_and_references_are_private():
    # Generated data may remove only the private author reference.
    generated=json.loads((ROOT/'platform/practice_quests.json').read_text(encoding='utf-8'))
    course=next(q for q in generated if q['id']=='practice_csv_report')
    source=copy.deepcopy(quest())
    for stage in source['stages']:stage.pop('reference_code')
    assert course==source


def test_project_check_transports_files_without_changing_verdicts():
    import subprocess
    subprocess.run(['node',str(ROOT/'platform/scripts/verify_project_files_unit.cjs')],check=True)


def test_three_stage_access_waits_for_libraries_and_evidence_is_owned(client,srv):
    register(client,name='report_owner')
    course=srv.pc.get_course('practice_csv_report')
    url='/api/project-courses/practice_csv_report'
    assert client.post(url+'/start').status_code==409
    for topic in course['unlock_topics']:solve_topic(client,srv,topic)
    assert client.post(url+'/start').status_code==200
    for index,stage in enumerate(course['stages']):
        if index:
            current=client.get(url).get_json()['course']['stages'][index]
            assert current['status']=='waiting'
            for topic in stage['requirements']['topics']:solve_topic(client,srv,topic)
        endpoint=url+'/stages/'+stage['id']
        before=(client.get('/api/wallet').get_json()['drops'],client.get('/api/progress').get_json()['xp'])
        assert client.post(endpoint+'/complete',json={'check':{'ok':True}}).status_code==400
        evidence={'checked':[item['id'] for item in stage['local_practice']['checklist']],
                  'reflection':'Учебные данные и два результата сверил с независимыми итогами.'}
        response=client.post(endpoint+'/complete',json={'code':'# own saved work','check':{'ok':True},'local_evidence':evidence})
        assert response.status_code==200,response.get_json()
        assert response.get_json()['stage']['local_evidence']==evidence
        assert before==(client.get('/api/wallet').get_json()['drops'],client.get('/api/progress').get_json()['xp'])
    assert client.get(url).get_json()['course']['status']=='completed'
    client.post('/api/logout');register(client,name='report_other')
    assert all(not s['local_evidence'] for s in client.get(url).get_json()['course']['stages'])
